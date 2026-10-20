"""Excel 紀錄：往下加、當日總熱量（每人各算各的）、被鎖住時重試、防護 4。"""

import os
from datetime import datetime

import pytest
from openpyxl import load_workbook

from foodbot.records import HEADERS, RecordBook, RecordError, Row, neutralize_formula

DAY1 = datetime(2026, 10, 20, 12, 0, 0)
DAY2 = datetime(2026, 10, 21, 8, 0, 0)


def row(uid=1, name="小明", kcal=500.0, when=DAY1, desc="便當", ing="白飯 200公克/280大卡", photo="photos/a.jpg", weight=300.0):
    return Row(when=when, uploader_name=name, uploader_id=uid, description=desc, ingredients=ing,
               weight_g=weight if kcal is not None else None, kcal=kcal, photo_rel_path=photo)


def make_book(tmp_path, **kw):
    return RecordBook(tmp_path / "records.xlsx", sleep=lambda s: None, **kw)


def rows_of(path):
    return list(load_workbook(path).active.iter_rows(values_only=True))


def test_creates_file_with_header_then_appends_without_overwriting(tmp_path):
    book = make_book(tmp_path)
    book.append(row(kcal=500))
    book.append(row(kcal=300))
    rows = rows_of(book.path)
    assert list(rows[0]) == HEADERS
    assert len(rows) == 3  # 表頭 + 兩列，第一列沒被覆蓋
    assert rows[1][5] == 500 and rows[2][5] == 300
    assert isinstance(rows[1][4], (int, float))  # 重量/熱量是真的數字


def test_daily_total_accumulates_and_is_written_in_row(tmp_path):
    book = make_book(tmp_path)
    assert book.append(row(kcal=500)) == 500
    assert book.append(row(kcal=300)) == 800
    assert [r[6] for r in rows_of(book.path)[1:]] == [500, 800]


def test_daily_total_is_per_user_and_per_day(tmp_path):
    book = make_book(tmp_path)
    book.append(row(uid=1, kcal=500))
    assert book.append(row(uid=2, name="小華", kcal=200)) == 200  # 小華從 0 開始
    assert book.append(row(uid=1, kcal=100, when=DAY2)) == 100  # 隔天重新計算
    assert book.append(row(uid=1, kcal=100)) == 600


def test_not_food_row_does_not_increase_total(tmp_path):
    book = make_book(tmp_path)
    book.append(row(kcal=500))
    total = book.append(row(kcal=None, desc="不是食物", ing="馬克杯"))
    assert total == 500
    last = rows_of(book.path)[-1]
    assert last[2] == "不是食物" and last[3] == "馬克杯" and last[4] is None and last[5] is None and last[6] == 500


def test_totals_are_rebuilt_from_excel_after_restart(tmp_path):
    book = make_book(tmp_path)
    book.append(row(uid=1, kcal=500))
    book.append(row(uid=2, name="小華", kcal=200))
    book.append(row(uid=1, kcal=None, desc="不是食物", ing="貓"))
    fresh = make_book(tmp_path)
    fresh.load_totals()
    assert fresh.today_total(1, DAY1.date()) == 500
    assert fresh.today_total(2, DAY1.date()) == 200
    assert fresh.append(row(uid=1, kcal=100)) == 600


def test_uploader_name_with_parentheses_still_parses(tmp_path):
    book = make_book(tmp_path)
    book.append(row(uid=42, name="小明 (本人)", kcal=100))
    fresh = make_book(tmp_path)
    fresh.load_totals()
    assert fresh.today_total(42, DAY1.date()) == 100


def test_locked_file_is_retried_then_succeeds(tmp_path, monkeypatch, capsys):
    book = make_book(tmp_path)
    real_replace = os.replace
    calls = {"n": 0}

    def flaky(src, dst):
        calls["n"] += 1
        if calls["n"] <= 2:
            raise PermissionError("file is open in Excel")
        return real_replace(src, dst)

    monkeypatch.setattr(os, "replace", flaky)
    assert book.append(row(kcal=500)) == 500
    assert calls["n"] == 3
    assert "警告" in capsys.readouterr().out  # 終端機有印警告
    assert len(rows_of(book.path)) == 2
    assert not (tmp_path / "records.xlsx.tmp").exists()


def test_locked_forever_raises_record_error_and_keeps_old_data(tmp_path, monkeypatch):
    book = make_book(tmp_path, retries=3)
    book.append(row(kcal=500))

    def always_locked(src, dst):
        raise PermissionError("locked")

    monkeypatch.setattr(os, "replace", always_locked)
    with pytest.raises(RecordError):
        book.append(row(kcal=300))
    monkeypatch.undo()
    assert len(rows_of(book.path)) == 2  # 舊資料完好
    assert book.today_total(1, DAY1.date()) == 500  # 失敗的那筆沒有計入
    assert not (tmp_path / "records.xlsx.tmp").exists()


# ---- 防護 4 ----
@pytest.mark.parametrize("text", ["=1+1", "+1", "-1", "@SUM(A1)", '=HYPERLINK("http://x","click")', "\t=1", "\r\n=1"])
def test_formula_like_text_is_neutralized(text):
    assert neutralize_formula(text).startswith("'")


def test_normal_text_is_unchanged():
    assert neutralize_formula("雞腿便當") == "雞腿便當"


def test_formula_text_is_stored_as_plain_text_in_excel(tmp_path):
    book = make_book(tmp_path)
    book.append(row(name="=1+1", desc="=2+2", ing='@SUM(1)', photo="photos/a.jpg"))
    ws = load_workbook(book.path).active
    for col in (2, 3, 4):  # 上傳者、食物敘述、成份
        cell = ws.cell(row=2, column=col)
        assert cell.data_type == "s", f"第 {col} 欄被當成公式了：{cell.value!r}"
        assert cell.value.startswith("'")


def test_illegal_control_characters_do_not_crash(tmp_path):
    book = make_book(tmp_path)
    book.append(row(desc="壞\x00字\x07元"))
    assert rows_of(book.path)[1][2] == "壞字元"
