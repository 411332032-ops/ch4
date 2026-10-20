"""寫 Excel 紀錄、算「當日已攝取熱量」。

包含：
  防護 4：寫進 Excel 的文字若以 = + - @ 開頭，處理成純文字

「一直往下加」的做法：開啟既有檔 → 在最底下加一列 → 先存成暫存檔 → 再原子取代原檔。
Excel 正開著（Windows 會鎖檔）時，印出警告並重試，不會讓程式當掉。
"""

from __future__ import annotations

import logging
import os
import re
import time
from collections import defaultdict
from collections.abc import Callable
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE

from .config import CELL_MAX_CHARS, EXCEL_RETRIES, EXCEL_RETRY_DELAY_SEC

log = logging.getLogger("foodbot.records")

HEADERS = [
    "日期時間",
    "上傳者",
    "食物敘述",
    "成份",
    "重量(公克)",
    "熱量(大卡)",
    "當日已攝取熱量(大卡)",
    "照片存放位置",
]
COL_WHEN, COL_UPLOADER, COL_KCAL = 0, 1, 5  # 欄位位置（從 0 算起）
_UPLOADER_ID = re.compile(r"\((\d+)\)\s*$")


# ----------------------------------------------------------------------------
# 沒有防護 4 時的預設：文字原樣寫進 Excel（開頭是 = 就會變成公式）
# ----------------------------------------------------------------------------
def neutralize_formula(text: str) -> str:
    return text


# ==== 防護 4：Excel 公式注入處理　開始 ====
# 威脅：文字以 = + - @ 開頭時，Excel 會把它當成公式執行
#       （例如顯示名稱叫 =HYPERLINK("惡意網址","點我")）。
# 做法：開頭是這些字元（或 Tab、換行）時，前面加一個單引號，讓它變成純文字。
_FORMULA_STARTS = ("=", "+", "-", "@", "\t", "\r", "\n")


def neutralize_formula(text: str) -> str:  # noqa: F811
    return f"'{text}" if text.startswith(_FORMULA_STARTS) else text
# ==== 防護 4 結束 ====


def clean_text(text: str) -> str:
    """Excel 不允許的控制字元要移除（否則 openpyxl 會丟例外），並限制單格長度。"""
    return ILLEGAL_CHARACTERS_RE.sub("", text)[:CELL_MAX_CHARS]


def text_cell(text: str) -> str:
    return neutralize_formula(clean_text(text))


class RecordError(Exception):
    """Excel 寫入失敗（重試後仍失敗）。"""


@dataclass(frozen=True)
class Row:
    when: datetime
    uploader_name: str
    uploader_id: int
    description: str  # 食物敘述（不是食物時為「不是食物」）
    ingredients: str  # 成份（不是食物時為判定出來的東西）
    weight_g: float | None
    kcal: float | None
    photo_rel_path: str  # 相對於 DATA_DIR 的路徑


class RecordBook:
    def __init__(
        self,
        path: Path,
        retries: int = EXCEL_RETRIES,
        retry_delay: float = EXCEL_RETRY_DELAY_SEC,
        sleep: Callable[[float], None] = time.sleep,
    ) -> None:
        self.path = path
        self._retries = retries
        self._delay = retry_delay
        self._sleep = sleep
        # {(使用者ID, "YYYY-MM-DD"): 當日熱量加總}。Excel 才是真正的紀錄，這只是快取。
        self._totals: dict[tuple[int, str], float] = defaultdict(float)

    # ---- 當日總熱量 ----
    def today_total(self, user_id: int, day: date) -> float:
        return self._totals[(user_id, day.isoformat())]

    def load_totals(self) -> None:
        """啟動時從 Excel 重建每人每天的熱量加總。"""
        self._totals.clear()
        if not self.path.exists():
            return
        self._with_retry(self._read_totals)

    def _read_totals(self) -> None:
        wb = load_workbook(self.path, read_only=True)
        try:
            for row in wb.active.iter_rows(min_row=2, values_only=True):
                if not row or len(row) <= COL_KCAL:
                    continue
                kcal = row[COL_KCAL]
                if isinstance(kcal, bool) or not isinstance(kcal, (int, float)):
                    continue
                match = _UPLOADER_ID.search(str(row[COL_UPLOADER] or ""))
                when = str(row[COL_WHEN] or "")[:10]
                if match and when:
                    self._totals[(int(match.group(1)), when)] += float(kcal)
        finally:
            wb.close()

    # ---- 新增一列 ----
    def append(self, row: Row) -> float:
        """在 Excel 最底下加一列，回傳該使用者當日（含這一餐）的熱量總數。"""
        key = (row.uploader_id, row.when.date().isoformat())
        new_total = self._totals[key] + (row.kcal or 0)
        self._with_retry(lambda: self._write_row(row, new_total))
        self._totals[key] = new_total  # 寫成功之後才更新快取
        return new_total

    def _write_row(self, row: Row, daily_total: float) -> None:
        self.path.parent.mkdir(parents=True, exist_ok=True)
        if self.path.exists():
            wb = load_workbook(self.path)
            ws = wb.active
        else:
            wb = Workbook()
            ws = wb.active
            ws.title = "紀錄"
            ws.append(HEADERS)
            for letter, width in zip("ABCDEFGH", (20, 28, 20, 50, 12, 12, 20, 55)):
                ws.column_dimensions[letter].width = width

        ws.append(
            [
                text_cell(row.when.strftime("%Y-%m-%d %H:%M:%S")),
                text_cell(f"{row.uploader_name} ({row.uploader_id})"),
                text_cell(row.description),
                text_cell(row.ingredients),
                None if row.weight_g is None else round(row.weight_g, 1),
                None if row.kcal is None else round(row.kcal, 1),
                round(daily_total, 1),
                text_cell(row.photo_rel_path),
            ]
        )

        tmp = self.path.with_name(self.path.name + ".tmp")
        try:
            wb.save(tmp)
            os.replace(tmp, self.path)  # 原子取代；Excel 開著時這裡會丟 PermissionError
        finally:
            if tmp.exists():
                try:
                    tmp.unlink()
                except OSError:
                    pass

    # ---- 被鎖住時重試 ----
    def _with_retry(self, action: Callable[[], None]) -> None:
        for attempt in range(1, self._retries + 1):
            try:
                action()
                return
            except PermissionError as exc:
                log.warning(
                    "⚠️ 無法寫入 %s（可能正被 Excel 開著）。第 %d/%d 次嘗試：%s",
                    self.path.name, attempt, self._retries, exc,
                )
                print(f"⚠️ 警告：{self.path.name} 無法寫入，請關閉 Excel。第 {attempt}/{self._retries} 次重試中…")
                if attempt < self._retries:
                    self._sleep(self._delay)
        raise RecordError(f"{self.path.name} 一直無法寫入（是不是還開著？）")
